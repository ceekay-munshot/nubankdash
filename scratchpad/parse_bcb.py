#!/usr/bin/env python3
"""Regenerate src/data/bcb.ts from the BCB credit-data workbook.

Setup (once):
    python3 -m pip install -r scratchpad/requirements.txt

Usage:
    python3 scratchpad/parse_bcb.py "path/to/<MM.DD.YYYY> - Bz Central Bank data.xlsx" [--check]

What the workbook is: a monthly pull of Banco Central do Brasil (SGS) household
credit-quality ratios. It is a companion to the quarterly Nu_Brazil workbook
(scratchpad/parse.py) and is now the dashboard's ONLY source for system-wide NPL
by product - it is monthly, so it runs a quarter or more ahead of the copy of
the same series the Nu_Brazil workbook carries at quarter ends.

Layout (both sheets identical):
  "NPL"        - % of the book 90+ days overdue
  "Early NPL " - % of the book 15-90 days overdue (the leading gauge)
Each sheet holds two side-by-side blocks, each with its own "Data" (date) column:
  left  - Household, Unsecured personal loans, Payroll loans, Credit card loans
  right - Payroll loans again, then its three borrower segments: private-sector
          workers, public employees, INSS retirees & pensioners
Row 5 carries the labels and row 4 the SGS series code above each one. Columns
are located by those labels, never by position, so an inserted column moves the
data rather than silently mislabelling it.

Method:
  * Values are fractions in the sheet; emitted as % to 2dp. The BCB publishes
    two decimals of a percent, so this is lossless.
  * The period axis is a contiguous calendar-month range. A month the sheet
    skips is emitted as null rather than closed up, so "12 readings back" is
    always exactly a year in the dashboard's maths.
  * "Payroll loans" appears in both blocks. The left (product) block is the
    canonical copy for every chart, so the payroll line reads identically in the
    product chart and the segment chart; the right-hand copy is only
    cross-checked, and disagreements are printed.
  * "Household" is BCB's ALL-household total (series 21084 / 21005), which
    includes mortgages, vehicle and rural credit. It is a broader series than
    the household line the Nu_Brazil workbook carried, so it is emitted as
    "All household loans" to keep the two from being read as one.

The workbook is trusted as uploaded: `--check` (and a normal run) prints what
the plausibility scan notices, for a human to read, but nothing here ever
refuses to write.
"""
from __future__ import annotations

import datetime as dt
import json
import re
import sys
from pathlib import Path

# Share the edition-stamp logic with the Nu_Brazil regenerator. Inserted
# explicitly so this also works under `python3 -I`, which drops the script's
# own directory from sys.path.
sys.path.insert(0, str(Path(__file__).resolve().parent))
from parse import js, num, workbook_edition  # noqa: E402

import openpyxl  # noqa: E402  (parse.py already explained how to install it)

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "src" / "data" / "bcb.ts"

SHEETS = {"npl": "NPL", "early": "Early NPL"}
LABEL_ROW = 5
CODE_ROW = 4
FIRST_DATA_ROW = 6
DATE_LABEL = "data"

# Sheet label -> dashboard series name. "Personal loans" keeps the name the
# dashboard has always used for this exact series (BCB "não consignado").
LABELS = {
    "credit card loans": "Credit card loans",
    "unsecured personal loans": "Personal loans",
    "payroll loans": "Payroll loans",
    "household": "All household loans",
    "private payroll loans": "Private-sector payroll",
    "public employee payroll loans": "Public-sector payroll",
    "inss payroll loans": "INSS payroll",
}
PRODUCT_ORDER = ["Credit card loans", "Personal loans", "Payroll loans", "All household loans"]
SEGMENT_ORDER = ["Private-sector payroll", "Public-sector payroll", "INSS payroll", "Payroll loans"]

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]


def month_label(y, m):
    return f"{MONTHS[m - 1]} {y}"


def find_sheet(wb, name):
    """Sheet titles in this workbook carry stray trailing spaces ("Early NPL ")."""
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    raise SystemExit(f'sheet "{name}" not found; workbook has {[ws.title for ws in wb.worksheets]}')


def sgs_code(v):
    m = re.match(r"\s*(\d{4,6})\b", str(v)) if v is not None else None
    return int(m.group(1)) if m else None


def read_blocks(ws):
    """-> list of blocks: {"date_col": c, "cols": {series name: (col, sgs code)}}.

    A block starts at a "Data" label and runs until the next one.
    """
    blocks = []
    for c in range(1, ws.max_column + 1):
        raw = ws.cell(LABEL_ROW, c).value
        if raw is None:
            continue
        label = str(raw).strip().lower()
        if label == DATE_LABEL:
            blocks.append({"date_col": c, "cols": {}})
            continue
        name = LABELS.get(label)
        if name is None or not blocks:
            print(f'  note: {ws.title.strip()} column {c} label "{raw}" is not wired')
            continue
        blocks[-1]["cols"][name] = (c, sgs_code(ws.cell(CODE_ROW, c).value))
    return blocks


def read_block(ws, block):
    """-> {(year, month): {series: value-in-% or None}}"""
    out = {}
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        d = ws.cell(r, block["date_col"]).value
        if not isinstance(d, (dt.date, dt.datetime)):
            continue
        key = (d.year, d.month)
        if key in out:
            print(f"  note: {ws.title.strip()} row {r} repeats {month_label(*key)}; later row wins")
        out[key] = {}
        for name, (c, _) in block["cols"].items():
            v = num(ws.cell(r, c).value)
            out[key][name] = None if v is None else round(100 * v, 2)
    return out


def month_axis(keys):
    """Contiguous (year, month) range covering every key."""
    if not keys:
        return []
    (y, m), end = min(keys), max(keys)
    axis = []
    while (y, m) <= end:
        axis.append((y, m))
        y, m = (y + 1, 1) if m == 12 else (y, m + 1)
    return axis


def to_chart(rows, order):
    axis = month_axis(list(rows))
    series = {n: [rows.get(k, {}).get(n) for k in axis] for n in order}
    # Drop trailing months nobody has filled in yet.
    while axis and all(series[n][-1] is None for n in order):
        axis.pop()
        for n in order:
            series[n].pop()
    return {"periods": [month_label(*k) for k in axis], "rows": series}


def excursions(name, periods, values):
    """FYI only: a reading far off on the SAME side of both neighbours.

    Monthly arrears ratios can genuinely jump for a month (a state payroll
    paid late shows up in public-sector early NPL, for instance), so this is
    printed for a human to look at, never acted on.
    """
    found = []
    for i in range(1, len(values) - 1):
        p, cur, n = values[i - 1], values[i], values[i + 1]
        if None in (p, cur, n) or min(p, n) <= 0:
            continue
        if (cur / p > 1.35 and cur / n > 1.35) or (cur / p < 0.74 and cur / n < 0.74):
            found.append(f"{name:24s} {periods[i]}: {p:.2f} -> {cur:.2f} -> {n:.2f}")
    return found


def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    charts, codes, notes = {}, {}, []
    for key, title in SHEETS.items():
        ws = find_sheet(wb, title)
        blocks = read_blocks(ws)
        product = next((b for b in blocks if "Credit card loans" in b["cols"]), None)
        segment = next((b for b in blocks if "Private-sector payroll" in b["cols"]), None)
        if product is None or segment is None:
            raise SystemExit(f"{title}: could not find both the product and the payroll-segment blocks")
        prod_rows, seg_rows = read_block(ws, product), read_block(ws, segment)

        # Cross-check the duplicated payroll column, then make the product
        # block's copy canonical for both charts.
        for k in sorted(set(prod_rows) | set(seg_rows)):
            a = prod_rows.get(k, {}).get("Payroll loans")
            b = seg_rows.get(k, {}).get("Payroll loans")
            if a is not None and b is not None and a != b:
                notes.append(f"{title}: payroll {month_label(*k)} reads {a:.2f} in the product block, {b:.2f} in the segment block - product block used")
            seg_rows.setdefault(k, {})["Payroll loans"] = a

        charts[key] = to_chart(prod_rows, PRODUCT_ORDER)
        charts[key + "Payroll"] = to_chart(seg_rows, SEGMENT_ORDER)
        for block in (product, segment):
            for name, (_, code) in block["cols"].items():
                codes.setdefault(name, {})[key] = codes.get(name, {}).get(key) or code
    return charts, codes, notes


# --- emit -------------------------------------------------------------------

def chart_ts(name, data, order, doc):
    lines = [f"/** {doc} */", f"export const {name}: ChartData = {{"]
    lines.append("  periods: [" + ", ".join(json.dumps(p) for p in data["periods"]) + "],")
    lines.append("  series: [")
    for key in order:
        vals = ", ".join(js(v) for v in data["rows"][key])
        lines.append(f"    {{ name: {json.dumps(key)}, values: [{vals}] }},")
    lines.append("  ],")
    lines.append("};")
    return "\n".join(lines)


HEADER = '''// AUTO-GENERATED from the BCB credit-data workbook ({edition}) - sheets "NPL"
// and "Early NPL". Do not edit by hand; regenerate with
// `python3 scratchpad/parse_bcb.py <workbook.xlsx>`.
//
// Monthly Banco Central do Brasil (SGS) household credit-quality ratios, %:
//  - systemNpl, payrollNpl            - share of the book 90+ days overdue
//  - systemEarlyNpl, payrollEarlyNpl  - share 15-90 days overdue (leading gauge)
// "Personal loans" = non-payroll ("nao consignado") personal credit.
// "All household loans" = BCB's all-household total, mortgages, vehicle and
// rural credit included - a broader series than the household line the
// Nu_Brazil workbook carried, hence the different name.
// "Payroll loans" is the same series in every chart (the workbook prints it
// twice; the product block's copy is used throughout).
// Months the sheet skips are null, so 12 readings back is always one year.
import type { ChartData } from "./datasets";

/** Workbook edition these series came from, stamped at generation time. */
export const BCB_WORKBOOK = "Bz Central Bank data.xlsx ({edition})";

/** BCB SGS series code behind each line, per measure, as printed in the workbook. */
export const BCB_CODES: Record<string, { npl: number | null; early: number | null }> = {
{codes}
};
'''


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    path = args[0]
    edition = workbook_edition(path)
    charts, codes, notes = parse(path)

    fyi = []
    for key in ("npl", "early", "nplPayroll", "earlyPayroll"):
        c = charts[key]
        for name, vals in c["rows"].items():
            if key.endswith("Payroll") and name == "Payroll loans":
                continue  # already scanned in the product chart
            fyi.extend(excursions(name, c["periods"], vals))

    def report():
        for n in notes:
            print(f"  note: {n}")
        print(f"plausibility scan: {len(fyi)} one-month excursion(s) - FYI, the series are written as published")
        for f in fyi:
            print(f"  {f}")

    if "--check" in sys.argv:
        print(f"workbook edition: {edition}")
        for key, c in charts.items():
            print(f"{key}: {c['periods'][0]} - {c['periods'][-1]} ({len(c['periods'])} months)")
            for name, vals in c["rows"].items():
                real = [v for v in vals if v is not None]
                last = next((v for v in reversed(vals) if v is not None), None)
                rec = " (record)" if last is not None and last >= max(real) else ""
                print(f"  {name:24s} latest {last:6.2f}{rec}  range {min(real):.2f}-{max(real):.2f}")
        report()
        return 0

    report()
    codes_ts = ",\n".join(
        f"  {json.dumps(n)}: {{ npl: {js(codes.get(n, {}).get('npl'))}, early: {js(codes.get(n, {}).get('early'))} }}"
        for n in PRODUCT_ORDER + SEGMENT_ORDER[:3]
    )
    parts = [HEADER.replace("{edition}", edition).replace("{codes}", codes_ts)]
    parts.append(chart_ts("systemNpl", charts["npl"], PRODUCT_ORDER,
                          "90-day NPL by product, system-wide (% of each product's book)."))
    parts.append(chart_ts("systemEarlyNpl", charts["early"], PRODUCT_ORDER,
                          "Early (15-90 day) NPL by product, system-wide (%)."))
    parts.append(chart_ts("payrollNpl", charts["nplPayroll"], SEGMENT_ORDER,
                          "90-day NPL on payroll loans, split by who the borrower is (%)."))
    parts.append(chart_ts("payrollEarlyNpl", charts["earlyPayroll"], SEGMENT_ORDER,
                          "Early (15-90 day) NPL on payroll loans, by borrower segment (%)."))
    OUT.write_text("\n\n".join(parts) + "\n", encoding="utf-8")
    print(f"wrote {OUT} (workbook edition {edition}, through {charts['npl']['periods'][-1]})")
    return 0


if __name__ == "__main__":
    sys.exit(main() or 0)
